"""
LPS-Beacon - Resource availability and matching script
=======================================================

Matches the resources rolling off the Project Resource Plan (supply, "PRP") to the Quote Line positions that are
starting (demand, "QL"). The standardising and matching logic is the same as the backend matching Lambda
(LPS-Beacon_Backend/lambda-forecastpipeline/matching_lambda.py), which produced matching_result_new.xlsx.

Inputs
    1. Project Resource Plan  - one row per resource per fiscal year, man-days (MDs) per month
    2. Quote Line             - one row per quoted position per fiscal year, MDs per month
    3. Position Name Mapping  - a "Standard Role" list (the target roles) and/or
                                Position Name -> Standard Role examples

    File types: .xlsx .xlsm .xltx .xltm .xls .xlt .xlsb .ods .csv .tsv (.txt files are not accepted). The file
    content decides how it is read (a CSV saved as .xlsx still works). Anything else - PDF, Word, password-protected, HTML saved as .xls -
    stops with an "ERROR: ..." message. (.xls needs xlrd, .xlsb needs pyxlsb, .ods needs odfpy.)
    Sheets: any names. The sheet that has the needed columns is used; the header row may be below title rows.
    Headers: matched by spelling variants (see COLUMN_ALIASES), e.g. "Role"/"Job Title" -> Position Name,
    "FY" -> Fiscal Year, "April"/"Apr-26"/"Apr MDs" -> Apr, "Band" -> Grade, "Skills" -> Skill.
    Values: Fiscal Year as FY26 / FY2026 / 2026 / 2026-27; Grade as 6 / G6 / Grade 6.

Outputs (default: the "Output Files" folder next to this script)
    Output Files/Latest/                         the four files from the newest run (replaced every run)
    Output Files/Historical Matches/<run time>/  a copy of every run, e.g. Historical Matches/2026-10-20_143005
    With --out-dir the four files are written straight into that folder instead (used by automation/ and aws_lambda/).
    matching_result_new.xlsx   Summary, Matching, Matched, Unmatched Resources, Unmatched Projects
    Clean Resource Plan.xlsx   the resource plan rows with their standardized position and dates
    Clean Quoteline.xlsx       the quote line rows with their standardized position and dates
    Unmatched Resources.xlsx   resources free in the window that no quote line could use
    All dates are shown as DD-MM-YY.

Date rules (all dates are the 1st of a month)
    Available From = the month AFTER the last month with man-days   (last MDs in Aug -> 01-09-26)
    Start Date     = the FIRST month with man-days                   (first MDs in Sep -> 01-09-26)
    A resource can fill a line when Available From <= Start Date, and it is not idle longer than --max-idle months.

Matching order (one resource per line; each pass is an optimal assignment)
    T1 standardized position + same grade
    T2 skill (used when position names do not match / are unmapped) + same grade
    T3 standardized position + grade +/- tolerance
    T4 skill + grade +/- tolerance

Usage (Command Prompt / PowerShell, from this folder)
    python matching_python_script.py
    python matching_python_script.py --as-of 2026-10-01
    python matching_python_script.py --resource-plan "Edited_Files/sample_rw_rp_edited.xlsx"
                                     --quoteline "Edited_Files/sample_rw_quoteline(edited).csv"

Needs: Python 3.10+, pandas, numpy, scipy, scikit-learn, openpyxl (+ xlrd, pyxlsb, odfpy for .xls/.xlsb/.ods)
    pip install pandas numpy scipy scikit-learn openpyxl xlrd pyxlsb odfpy
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import shutil
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linear_sum_assignment
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# =========================================================================== SETTINGS (all can be changed on the command line)
HERE = Path(__file__).resolve().parent
DEFAULT_RESOURCE_PLAN = HERE / "sample_rw_rp(based on  Project Resource Plan)(Query1).csv"
DEFAULT_QUOTELINE = HERE / "sample_rw_rp(based on Quoteline)(Query1).csv"
DEFAULT_MAPPING = HERE / "Position Name Mapping (1).xlsx"

# Where results go when --out-dir is not given: Latest/ holds the newest run, Historical Matches/<run time>/ every run
OUTPUT_ROOT = HERE / "Output Files"
LATEST_DIR = "Latest"
HISTORY_DIR = "Historical Matches"
RUN_ID_FORMAT = "%Y-%m-%d_%H%M%S"

# The four output files
OUT_MATCHING = "matching_result_new.xlsx"
OUT_CLEAN_RP = "Clean Resource Plan.xlsx"
OUT_CLEAN_QL = "Clean Quoteline.xlsx"
OUT_UNMATCHED = "Unmatched Resources.xlsx"

DATE_FORMAT = "dd-mm-yy"            # how every date cell is shown in Excel

FY_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
MONTH_NUM = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
GRADE_BAND = {3: "B1", 4: "B1", 5: "B2", 6: "B2", 7: "B3", 8: "B4", 9: "B5", 10: "B5", 11: "B5", 12: "B5"}
BIG = 1e6
UNMAPPED = "UNMAPPED - FOR REVIEW"

# Labels typed in the mapping sheet that are not spelled like the "Standard Role" list -> canonical standard role
LABEL_ALIASES = {"service delivery manager": "SDM", "service delivery director": "SDD",
                 "solution architect": "Technical Architect", "operation manager": "SDM"}
# Only these roles read a trailing number as a level ("Cloud Engineer 1" == "Cloud Engineer L1").
# For other roles a trailing number is an instance count (ITSO 2, DC Ops 7).
LEVEL_DIGIT_BASES = {"Cloud Engineer"}

# Ordered rules: first match wins. Left side must be a base name that exists in the "Standard Role" sheet.
ROLE_RULES = [
    ("IT Security Officer", r"\bitso\b|\bsito\b|it security officer|security officer|\biso engineer|^security\s*-\s*l\d"),
    ("IT Security Mgr", r"\bciso\b|security manager|security mgr"),
    ("SDD", r"\bsdd\b|service delivery director|delivery director"),
    ("Senior Project Manager", r"senior project manager|sr\.? project manager"),
    ("Application PM", r"application pm|application project manager"),
    ("Project Manager", r"project manager|project lead|\bpm\b|project mgr"),
    ("SDM", r"\bsdm\b|service delivery (manager|lead)|\bsdl\b|operations? manager|delivery manager|delivery lead|"
            r"dc manager|incident manager|command cent(re|er) manager|cloud technical manager|cloud shared service manager"),
    ("Service Desk Mgr", r"service ?desk\s*(manager|mgr)"),
    ("Service Desk Lead", r"service ?desk\s*(lead|supervisor)"),
    ("Service Desk Agent", r"service ?desk|help ?desk|response cent(re|er)|\bnoc\b"),
    ("Enterprise Architect", r"enterprise architect"),
    ("Application Architect", r"application architect"),
    ("Cloud Architect", r"cloud arc\w*"),
    ("Application Support Lead", r"app(lication)? (support )?lead"),
    ("Application Support Engineer", r"app(lication)? (support )?engineer|application support|app support"),
    ("Technical Architect", r"arc\w*t\w*ect|\barchitect\b|\bsa\b"),
    ("Security Engineer", r"security engineer|cyber ?ark|\bsoc\b|\biam\b|\bpam\b"),
    ("Cloud DBA", r"cloud dba"),
    ("DBA", r"\bdba\b|database"),
    ("Site Reliability Engineer", r"\bsre\b|site reliability"),
    ("FinOps Engineer", r"finops"),
    ("M365 Engineer", r"m365|microsoft 365|o365|office 365|exchange"),
    ("SCCM Engineer", r"sccm|intune|endpoint manager"),
    ("Cloud Engineer", r"cloud|\bgcc\b|\baws\b|azure|\bgcp\b"),
    ("Mobile/Desktop Engineer", r"desktop|mobile|\bdex\b|end ?user|workplace|tech refresh"),
    ("Server Engineer", r"wintel|windows|linux|unix|server|system(s)? engineer|\bsys eng|middleware|vmware|virtuali[sz]ation"),
    ("Network Engineer", r"network|firewall|\bf5\b|load balanc"),
    ("Backup/Storage Engineer", r"storage|back-?up"),
    ("Tools Engineer (Observability)", r"observab|monitoring|tools engineer|\bapm\b"),
    ("IT Service Mgt Engineer", r"itsm|service management|change engineer|servicenow"),
    ("DCOps", r"dc ?op|data cent(re|er)|it ?/? ?(sr )?admin|asset (admin|officer|manager)|\badmin\b|\bdc\b"),
    ("AI Developer", r"\bai\b|machine learning|genai|\bllm\b"),
    ("Frontend Developer", r"front ?end|ui developer"),
    ("Data Engineer", r"data engineer|\betl\b"),
    ("Data Analyst", r"data analyst|\bmis\b"),
    ("Functional Analyst", r"functional analyst"),
    ("Business Analyst", r"business analyst|\bba\b"),
    ("Tester/QA", r"tester|\bqa\b|\buat\b|testing"),
    ("Full Stack Engineer", r"full ?stack|developer|\bjava\b|programmer"),
]

# Skill keyword families (Skill column = ';' separated competencies / certifications)
SKILL_FAMILIES = {
    "SECURITY": ["cyber", "security", "cissp", "cism", "cisa", "siem", "dlp", "edr", "vulnerab", "ccsp", "audit and compliance"],
    "NETWORK": ["network", "cisco", "ccna", "ccnp", "f5 ", "juniper", "checkpoint", "fortinet", "palo alto", "router", "switch"],
    "DATABASE": ["database", "oracle", "sql", "dba", "mysql", "postgres", "mongodb", "redis", "ocp"],
    "CLOUD": ["cloud", "aws", "azure", "gcp", "kubernetes", "docker", "devops", "terraform", "ci/cd"],
    "INFRA_OS": ["linux", "unix", "windows", "wintel", "active directory", "vmware", "vsphere", "storage", "backup",
                 "solaris", "data center", "netapp", "san "],
    "SERVICE_MGMT": ["itil", "servicenow", "itsm", "service desk", "incident", "service management", "remedy", "bmc"],
    "PROJECT_MGMT": ["project management", "pmp", "scrum", "agile", "pmbok", "primavera"],
    "WORKPLACE": ["desktop", "endpoint", "m365", "microsoft 365", "digital workplace", "outlook", "hardware", "end user"],
    "DATA_ANALYTICS": ["power bi", "analytics", "data analysis", "business intelligence", "rpa", "uipath"],
}
# skill family -> base role used ONLY when the position name gives no mapping (2nd-level check)
FAMILY_TO_BASE = {"SECURITY": "IT Security Officer", "NETWORK": "Network Engineer", "DATABASE": "DBA", "CLOUD": "Cloud Engineer",
                  "INFRA_OS": "Server Engineer", "SERVICE_MGMT": "IT Service Mgt Engineer", "PROJECT_MGMT": "Project Manager",
                  "WORKPLACE": "Mobile/Desktop Engineer", "DATA_ANALYTICS": "Data Analyst"}
BASE_TO_FAMILY = {"IT Security Officer": "SECURITY", "IT Security Mgr": "SECURITY", "Security Engineer": "SECURITY",
                  "Network Engineer": "NETWORK", "DBA": "DATABASE", "Cloud DBA": "DATABASE",
                  "Cloud Engineer": "CLOUD", "Cloud Architect": "CLOUD", "FinOps Engineer": "CLOUD", "Site Reliability Engineer": "CLOUD",
                  "Server Engineer": "INFRA_OS", "Backup/Storage Engineer": "INFRA_OS", "DCOps": "INFRA_OS",
                  "IT Service Mgt Engineer": "SERVICE_MGMT", "Service Desk Mgr": "SERVICE_MGMT", "Service Desk Lead": "SERVICE_MGMT",
                  "Service Desk Agent": "SERVICE_MGMT", "SDM": "SERVICE_MGMT", "SDD": "PROJECT_MGMT",
                  "Project Manager": "PROJECT_MGMT", "Senior Project Manager": "PROJECT_MGMT", "Application PM": "PROJECT_MGMT",
                  "Mobile/Desktop Engineer": "WORKPLACE", "M365 Engineer": "WORKPLACE", "SCCM Engineer": "WORKPLACE",
                  "Data Analyst": "DATA_ANALYTICS", "Data Engineer": "DATA_ANALYTICS", "Business Analyst": "DATA_ANALYTICS"}


# =========================================================================== EXTRACT - read the input files
# Accepted uploads. The file CONTENT decides how it is read, so a CSV saved as ".xlsx" or a workbook without an
# extension still works; anything that is not a spreadsheet stops the run with a clear message.
EXCEL_TYPES = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlt", ".xlsb", ".ods"}
TEXT_TYPES = {".csv", ".tsv", ".tab", ".prn", ""}
REJECTED_TYPES = {".txt"}           # plain text files are never read, whatever their content
ACCEPTED_TYPES = "Excel (.xlsx .xlsm .xltx .xltm .xls .xlt .xlsb), OpenDocument (.ods) or CSV/TSV"
ENGINE_PACKAGE = {"openpyxl": "openpyxl", "xlrd": "xlrd", "pyxlsb": "pyxlsb", "odf": "odfpy"}
HEADER_SCAN_ROWS = 30               # the header row may sit below title / blank rows

# Header aliases. Headers are compared lower-case with punctuation turned into spaces, and must match the whole
# pattern (so "Name" is the project name but "Resource Name" is the resource). An exact standard header always wins.
COLUMN_ALIASES = {
    "Position Name": r"position( name| title)?|position ?desc(ription)?|role( name| title)?|job( title| role| name)?|"
                     r"designation|title|quote ?line ?position",
    "Fiscal Year": r"fiscal( year| yr)?|fy|f y|financial (year|yr)|fin year|year",
    "Grade": r"grade|grades|band|job (grade|band)|grade (level|band)|career (level|grade)|staff grade",
    "Skill": r"skills?|skill ?sets?|competenc(y|ies)|certifications?|skills? (and|&)? ?certifications?",
    "Project ID": r"project( id| no| number| code| ref(erence)?)|proj( id| no)?|project",
    "Quote Line Number": r"quote ?line( (number|no|num|id|ref))?|ql( (number|no|id))?|line( (number|no|id))",
    "Resource": r"resource( name)?|employee( name)?|staff( name)?|consultant( name)?|assigned resource|person",
    "Project Name": r"project (name|title|description)|name",
    "Customer": r"customer( name)?|client( name)?|account( name)?|end customer( name)?",
    "Modified Date": r"modified( date| on)?|last modified( date| on)?|updated( date| on)?|last updated",
    "Created On": r"created( date| on)?|creation date|date created",
}
_MONTH_WORD = (r"(jan(uary)?|feb(ruary)?|mar(ch)?|apr(il)?|may|june?|july?|aug(ust)?|sep(t(ember)?)?|oct(ober)?|"
               r"nov(ember)?|dec(ember)?)")
MONTH_HEADER = re.compile(rf"^(fy ?\d{{2,4}} )?{_MONTH_WORD}( \d{{2,4}})?( (mds?|man ?days?|days|effort|fte))?$")
MONTH_NUM_HEADER = re.compile(r"^(\d{4}) (\d{1,2})( \d{1,2})?( 00 00 00)?$")      # 2026-04 / 2026-04-01 00:00:00
# Sheet-name hints, used only to choose between several sheets that all have the needed columns
SHEET_HINTS = {"resource plan": r"resource|plan|prp|supply|roll", "quote line": r"quote|\bql\b|demand|pipeline"}


def norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(h).lower()).strip()


def canon_column(h):
    """Standard column name for a header cell, or None when it is not one the script uses."""
    if isinstance(h, (pd.Timestamp, date)):                       # Excel stores "Apr-26" headers as real dates
        return pd.Timestamp(h).strftime("%b")
    s = norm_header(h)
    if not s or s.startswith("unnamed"):
        return None
    m = MONTH_HEADER.match(s)
    if m:
        return m.group(2)[:3].title()
    m = MONTH_NUM_HEADER.match(s)
    if m and 1 <= int(m.group(2)) <= 12:
        return list(MONTH_NUM)[int(m.group(2)) - 1]
    for canon, rx in COLUMN_ALIASES.items():
        if s == norm_header(canon) or re.fullmatch(rx, s):
            return canon
    return None


def fail(msg: str):
    sys.exit(f"ERROR: {msg}")


def warn(msg: str):
    print(f"WARNING: {msg}")


def detect_format(path: Path) -> str:
    """'openpyxl' / 'xlrd' / 'pyxlsb' / 'odf' (the pandas engine) or 'text'. Stops on anything else."""
    ext = path.suffix.lower()
    if ext in REJECTED_TYPES:
        fail(f"{path.name}: '{ext}' files are not accepted. Upload {ACCEPTED_TYPES}.")
    with open(path, "rb") as fh:
        head = fh.read(65536)
    if not head:
        fail(f"{path.name} is empty.")
    if head.startswith(b"PK\x03\x04"):                             # zip container: xlsx / xlsb / ods / docx / pptx ...
        try:
            with zipfile.ZipFile(path) as z:
                names = set(z.namelist())
                mime = z.read("mimetype").decode("ascii", "ignore") if "mimetype" in names else ""
        except zipfile.BadZipFile:
            fail(f"{path.name} is damaged (not a readable Excel workbook). Open it in Excel and save it again.")
        if "xl/workbook.xml" in names:
            return "openpyxl"
        if "xl/workbook.bin" in names:
            return "pyxlsb"
        if "spreadsheet" in mime:
            return "odf"
        kind = ("a Word document" if any(n.startswith("word/") for n in names) else
                "a PowerPoint file" if any(n.startswith("ppt/") for n in names) else
                "an OpenDocument text/presentation file" if mime else "a ZIP archive")
        fail(f"{path.name} is {kind}, not an Excel file. Upload {ACCEPTED_TYPES}.")
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):       # old Office format (.xls) - or an encrypted .xlsx
        if "EncryptedPackage".encode("utf-16-le") in path.read_bytes():
            fail(f"{path.name} is password-protected. Remove the password in Excel (File > Info > Protect Workbook) "
                 f"and upload it again.")
        if ext in (".doc", ".ppt", ".msg"):
            fail(f"{path.name} is not an Excel file. Upload {ACCEPTED_TYPES}.")
        return "xlrd"
    if head.startswith(b"%PDF"):
        fail(f"{path.name} is a PDF, not an Excel file. Upload {ACCEPTED_TYPES}.")
    text = decode_text(head, path.name, strict=False)
    if text is None:
        fail(f"{path.name} is not an Excel or CSV file (binary content). Upload {ACCEPTED_TYPES}.")
    if re.match(r"\s*<", text):
        if "urn:schemas-microsoft-com:office:spreadsheet" in text:
            fail(f"{path.name} is an 'XML Spreadsheet 2003' file. Open it in Excel and Save As Excel Workbook (.xlsx).")
        fail(f"{path.name} is a web page / XML file, not a real Excel file (some systems export HTML with an .xls name). "
             f"Open it in Excel and Save As Excel Workbook (.xlsx).")
    if ext not in TEXT_TYPES | EXCEL_TYPES:
        fail(f"{path.name}: '{ext}' files are not accepted. Upload {ACCEPTED_TYPES}.")
    return "text"


def decode_text(data: bytes, name: str, strict=True):
    """Text in UTF-8 / UTF-16 (Excel 'Unicode Text') / Windows-1252. None (or stop, when strict) for binary data."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "replace")
    if b"\x00" in data:
        if strict:
            fail(f"{name} is not a text/CSV file.")
        return None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def read_sheets(path: Path) -> dict[str, pd.DataFrame]:
    """Every sheet of the file as a raw grid of cells (the header row is found later). A CSV is one sheet."""
    path = Path(path)
    if not path.exists():
        fail(f"Input file not found: {path}")
    if path.is_dir():
        fail(f"{path} is a folder, not a file.")
    engine = detect_format(path)
    if engine == "text":
        text = decode_text(path.read_bytes(), path.name)
        sample = "\n".join(text.splitlines()[:50])
        try:
            sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            sep = "\t" if sample.count("\t") > sample.count(",") else ","
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=sep))
        width = max((len(r) for r in rows), default=0)
        grid = pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)
        grid = grid.replace(r"^\s*$", None, regex=True)
        return {path.stem: grid}
    try:
        return pd.read_excel(path, sheet_name=None, header=None, engine=engine)
    except ImportError:
        pkg = ENGINE_PACKAGE[engine]
        fail(f"{path.name} is a {path.suffix or engine} workbook; reading it needs the Python package '{pkg}' "
             f"(pip install {pkg}), or save the file as .xlsx.")
    except Exception as e:
        fail(f"Cannot open {path.name} as an Excel workbook ({type(e).__name__}: {e}). "
             f"Open it in Excel and save it again as .xlsx.")


# File-name patterns used to find an input that was not given on the command line. Each file is claimed by the first
# kind whose pattern it matches (mapping, then quote line, then resource plan); this script's own outputs are skipped.
INPUT_NAME_PATTERNS = {
    "mapping": r"mapping|standard ?role|role ?(map|list|lookup)",
    "quote line": r"quote|(?<![a-z])ql(?![a-z])",
    "resource plan": r"resource|(?<![a-z])rp(?![a-z])|prp|plan",
}
OUTPUT_NAME = r"^(~\$|clean |unmatched |matching_result)"


def locate_input(given, default: Path, kind: str) -> Path:
    """The file given on the command line; else the default file name; else the newest file in this script's folder
    (or the current folder) whose name looks like this kind of input, with an accepted file type."""
    if given:
        path = Path(given)
        if path.exists():
            return path
        near = [f.name for f in (path.parent if str(path.parent) not in ("", ".") else Path.cwd()).glob("*")
                if f.is_file() and f.suffix.lower() in EXCEL_TYPES | TEXT_TYPES][:15]
        fail(f"{kind} file not found: {path}" + (f". Files in that folder: {near}" if near else ""))
    if Path(default).exists():
        return Path(default)
    found = []
    for folder in dict.fromkeys([HERE, Path.cwd()]):
        for f in folder.glob("*"):
            name = f.name.lower()
            if not f.is_file() or f.suffix.lower() not in EXCEL_TYPES | TEXT_TYPES - {""} or re.search(OUTPUT_NAME, name):
                continue
            claimed = next((k for k, rx in INPUT_NAME_PATTERNS.items() if re.search(rx, name)), None)
            if claimed == kind:
                found.append(f)
    if not found:
        flag = {"mapping": "--mapping", "quote line": "--quoteline", "resource plan": "--resource-plan"}[kind]
        fail(f"no {kind} file found in {HERE} (looked for a name like {INPUT_NAME_PATTERNS[kind]!r}). "
             f"Pass it with {flag} \"<file>\".")
    found.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    others = f" (newest of {len(found)}: {[f.name for f in found[1:4]]} also match)" if len(found) > 1 else ""
    print(f"{kind}: using {found[0].name}{others}")
    return found[0]


def find_header(grid: pd.DataFrame, classify=None):
    """(row index, score): the row among the first rows with the most recognised column headers."""
    classify = classify or canon_column
    best = (0, -1)
    for i in range(min(HEADER_SCAN_ROWS, len(grid))):
        score = sum(classify(v) is not None for v in grid.iloc[i] if pd.notna(v))
        if score > best[1]:
            best = (i, score)
    return best


def frame_from_grid(grid: pd.DataFrame, header_row: int) -> pd.DataFrame:
    """Rows below the header row, headers tidied, empty rows/columns removed, numbers typed as numbers."""
    names, seen = [], {}
    for i, h in enumerate(grid.iloc[header_row]):
        if pd.isna(h):
            n = f"Column {i + 1}"
        elif isinstance(h, (pd.Timestamp, date)):
            n = pd.Timestamp(h).strftime("%b-%y")
        else:
            n = re.sub(r"\s+", " ", str(h)).strip() or f"Column {i + 1}"
        seen[n] = seen.get(n, 0) + 1
        names.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
    df = grid.iloc[header_row + 1:].copy()
    df.columns = names
    df = df.dropna(how="all")
    df = df.loc[:, [c for c in df.columns if not (c.startswith("Column ") and df[c].isna().all())]]
    df = df.reset_index(drop=True)
    for c in df.columns:                                          # "1,234" / "18" text -> numbers when the whole column is numeric
        if df[c].dtype == object:
            vals = df[c].dropna()
            if len(vals) and not vals.map(lambda v: isinstance(v, (pd.Timestamp, date))).any():
                num = pd.to_numeric(vals.astype(str).str.replace(r"(?<=\d),(?=\d{3}\b)", "", regex=True).str.strip(),
                                    errors="coerce")
                if num.notna().all() and not vals.astype(str).str.match(r"^0\d").any():   # keep IDs like 00123 as text
                    df[c] = pd.to_numeric(df[c].astype("string").str.replace(r"(?<=\d),(?=\d{3}\b)", "", regex=True),
                                          errors="coerce")
    return finite(df)


def finite(df: pd.DataFrame) -> pd.DataFrame:
    """'inf' / 'Infinity' / #DIV/0! style values -> empty. Infinity is not valid JSON, so it must never reach the
    Summary or the result files that the matching Lambda sends back to Power Automate."""
    return df.replace([np.inf, -np.inf], np.nan)


def standardise_headers(df: pd.DataFrame, where: str) -> pd.DataFrame:
    """Renames recognised headers to the standard names. An exact standard name is kept before any alias."""
    canon = {c: canon_column(c) for c in df.columns}
    taken = {c for c in df.columns if c in COLUMN_ALIASES or c in MONTH_NUM}
    rename = {}
    for c, k in canon.items():
        if not k or c == k or c in taken and c not in rename:
            continue
        if k in taken:
            if k in MONTH_NUM and df[c].notna().any():
                first = next(o for o in df.columns if (rename.get(o, o) == k))
                fail(f"{where}: the columns '{first}' and '{c}' both look like the month {k}. Keep one fiscal year "
                     f"per row (Apr..Mar) with a 'Fiscal Year' column.")
            continue
        rename[c] = k
        taken.add(k)
    return df.rename(columns=rename)


def parse_fiscal_year(v):
    """FY26, FY 26, FY2026, 2026, 2026/27, FY26-27, 26 -> 'FY26' (FY26 = Apr 2026 - Mar 2027). None if unreadable."""
    if pd.isna(v):
        return None
    if isinstance(v, (pd.Timestamp, date)):
        return None
    m = re.search(r"(\d{4}|\d{1,2})", str(v).replace(".0", "") if isinstance(v, float) else str(v))
    if not m:
        return None
    n = int(m.group(1))
    year = n if n >= 1000 else 2000 + n
    return f"FY{year % 100:02d}" if 2000 <= year <= 2099 else None


def parse_grade(v):
    """6, '6', 6.0, 'G6', 'Grade 6', 'GR-06', 'JG6' -> 6.0. Anything else -> NaN."""
    if pd.isna(v):
        return np.nan
    if isinstance(v, (int, float, np.number)):
        return float(v)
    m = re.fullmatch(r"\s*(?:g|gr|grade|jg|job grade)?\s*[-.:]?\s*(\d{1,2})(?:\.0+)?\s*", str(v), re.I)
    return float(m.group(1)) if m else np.nan


def read_table(path: Path, kind: str = "input") -> pd.DataFrame:
    """Reads a resource plan or quote line file: CSV or any Excel type, any sheet name, header row anywhere in the
    first rows, headers in the usual spellings. Stops with a clear message when the needed columns are not there."""
    path = Path(path)
    where = f"{kind} file '{path.name}'"
    sheets = read_sheets(path)
    if not sheets:
        fail(f"{where} has no sheets.")

    hint = SHEET_HINTS.get(kind, r"$^")
    candidates = []
    for name, grid in sheets.items():
        grid = grid.dropna(how="all").dropna(axis=1, how="all")   # index keeps the sheet's row numbers
        if grid.empty:
            continue
        row, score = find_header(grid)
        cols = {canon_column(v) for v in grid.iloc[row] if pd.notna(v)}
        has_req = {"Position Name", "Fiscal Year"} <= cols and bool(cols & set(FY_MONTHS))
        candidates.append(((has_req, bool(re.search(hint, str(name), re.I)), score, len(grid)), name, grid, row, cols))
    if not candidates:
        fail(f"{where} has no data in any sheet.")
    candidates.sort(key=lambda c: c[0], reverse=True)
    (has_req, _, _, _), sheet, grid, row, cols = candidates[0]
    if not has_req:
        missing = [c for c in ("Position Name", "Fiscal Year") if c not in cols]
        if not cols & set(FY_MONTHS):
            missing.append("month columns (Apr, May ... Mar)")
        headers = [str(v) for v in grid.iloc[row] if pd.notna(v)][:40]
        fail(f"{where}: no sheet has the columns the matching needs. Missing: {', '.join(missing)}. "
             f"Best sheet '{sheet}' has the headers: {headers}")
    if len(sheets) > 1:
        print(f"{kind}: using sheet '{sheet}' of {path.name} (header on row {grid.index[row] + 1})")

    df = standardise_headers(frame_from_grid(grid, row), where)
    if df.empty:
        fail(f"{where}: the sheet '{sheet}' has headers but no data rows.")
    if "Customer" not in df.columns and "End Customer" in df.columns:
        df["Customer"] = df["End Customer"]

    fy = df["Fiscal Year"].map(parse_fiscal_year)
    bad = fy.isna() & df["Fiscal Year"].notna()
    if fy.isna().all():
        fail(f"{where}: no 'Fiscal Year' value could be read (examples: {df['Fiscal Year'].dropna().astype(str).head(3).tolist()}). "
             f"Use values like FY26, FY2026 or 2026.")
    if bad.any():
        warn(f"{where}: {int(bad.sum())} rows skipped - Fiscal Year not readable "
             f"(e.g. {df.loc[bad, 'Fiscal Year'].astype(str).head(3).tolist()})")
    if (fy.isna() & ~bad).any():
        warn(f"{where}: {int((fy.isna() & ~bad).sum())} rows skipped - no Fiscal Year (totals or blank rows)")
    df = df[fy.notna()].copy()
    df["Fiscal Year"] = fy[fy.notna()].astype("string")

    if "Grade" in df.columns:
        grade = df["Grade"].map(parse_grade)
        if grade.isna().all() and df["Grade"].notna().any():
            warn(f"{where}: no Grade value is a number (e.g. {df['Grade'].dropna().astype(str).head(3).tolist()}) - "
                 f"grade matching is off for this file")
        df["Grade"] = grade
    else:
        warn(f"{where}: no Grade column - grade matching is off for this file")
        df["Grade"] = np.nan
    for c, label in (("Skill", "skill matching (T2/T4) only uses the position"), ("Project ID", "lines are told apart "
                     "by position and resource only"), ("Resource", "resources are told apart by line only")):
        if c not in df.columns:
            warn(f"{where}: no '{c}' column - {label}")
    months = [m for m in FY_MONTHS if m in df.columns]
    for m in months:
        df[m] = pd.to_numeric(df[m], errors="coerce")
    if not (df[months].fillna(0) > 0).any().any():
        fail(f"{where}: no row has man-days above 0 in the month columns {months} - check that they hold numbers.")
    return df


# =========================================================================== taxonomy (from the Position Name Mapping workbook)
class Taxonomy:
    def __init__(self, roles, redirects, examples):
        self.roles = roles                                  # canonical standard roles
        self.canon = {r.lower(): r for r in roles}
        self.redirects = redirects                          # "SR/IT Admin" -> "DCOps"
        self.parse, self.levels, self.role_of = {}, {}, {}
        for r in roles:
            m = re.match(r"^(.*?)\s+L(\d)\b\s*(\(.*\))?$", r)
            base, lvl = ((m.group(1) + (" " + m.group(3) if m.group(3) else "")).strip(), f"L{m.group(2)}") if m else (r, "")
            self.parse[r] = (base, lvl)
            self.role_of[(base, lvl)] = r
            if lvl:
                self.levels.setdefault(base, []).append(lvl)
        for b in self.levels:
            self.levels[b] = sorted(self.levels[b])
        self.examples = examples                            # norm_key(position name) -> standard role

    def role_name(self, base, level):
        """Standard-form name for base+level. Uses the Standard Role sheet string when it exists; otherwise builds it in the
        same convention (e.g. Cloud Engineer L3 when the sheet only lists L1/L2)."""
        r = self.role_of.get((base, level))
        if r or not level:
            return r or self.role_of.get((base, ""))
        m = re.match(r"^(.*?)(\s*\(.*\))?$", base)
        return f"{m.group(1)} {level}{m.group(2) or ''}"

    def canonical(self, label):
        raw = str(label).strip().lower()
        k = LABEL_ALIASES.get(raw, raw)
        k = self.redirects.get(k.lower(), k)
        return self.canon.get(str(k).lower()) or self.canon.get(raw)


def norm_key(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


STD_ROLE_HEADER = r"standard( role| roles| position| title)?|std( role| position)?|standardi[sz]ed( role| position)?|" \
                  r"target role|mapped( role| to)?|maps? to|role mapping|canonical role"
STD_SHEET_NAME = r"standard|\bstd\b|role ?list|^roles?$|taxonomy|master"
MAP_SHEET_NAME = r"mapping|\bmap\b|lookup|position"


def mapping_column(h):
    """'Standard Role' / 'Position Name' / None for a header of the mapping file (std-role spellings are tried first,
    because 'Role' alone means the position name in the data files)."""
    s = norm_header(h)
    if re.fullmatch(STD_ROLE_HEADER, s):
        return "Standard Role"
    return "Position Name" if canon_column(h) == "Position Name" else None


def load_taxonomy(path):
    """Reads the mapping file: any Excel type or CSV, any sheet names. Finds
       - the mapping sheet: a sheet with a Position Name column and a Standard Role column
       - the standard role sheet: a one-column role list (sheet name like 'Standard Role', 'Roles', 'Std Role List').
    With no role list (e.g. a CSV with only the mapping) the standard roles are the values of the Standard Role column."""
    path = Path(path)
    where = f"mapping file '{path.name}'"
    sheets = read_sheets(path)
    mapping, role_list = None, None                               # (preference, sheet name, frame)
    for name, grid in sheets.items():
        grid = grid.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
        if grid.empty:
            continue
        row, _ = find_header(grid, mapping_column)
        heads = {i: mapping_column(v) for i, v in enumerate(grid.iloc[row]) if pd.notna(v)}
        kinds = set(heads.values())
        name_s = str(name)
        if {"Position Name", "Standard Role"} <= kinds:
            pos_i = next(i for i, k in heads.items() if k == "Position Name")
            std_i = next(i for i, k in heads.items() if k == "Standard Role")
            frame = pd.DataFrame({"Position Name": grid.iloc[row + 1:, pos_i].values,
                                  "Standard Role": grid.iloc[row + 1:, std_i].values})
            pref = (bool(re.search(MAP_SHEET_NAME, name_s, re.I)), int(frame["Standard Role"].notna().sum()))
            if mapping is None or pref > mapping[0]:
                mapping = (pref, name_s, frame)
        elif "Position Name" not in kinds:
            # role list: the Standard Role column, else the first column of a sheet named like a role list
            if "Standard Role" in kinds:
                col, start = next(i for i, k in heads.items() if k == "Standard Role"), row + 1
            elif re.search(STD_SHEET_NAME, name_s, re.I) and grid.shape[1] <= 3:
                col, start = 0, 0
            else:
                continue
            values = grid.iloc[start:, col].dropna().astype(str).str.strip()
            values = values[values.ne("") & ~values.map(lambda v: mapping_column(v) == "Standard Role"
                                                         or bool(re.fullmatch(r"roles?( list)?", norm_header(v))))]
            pref = (bool(re.search(STD_SHEET_NAME, name_s, re.I)), len(values))
            if len(values) and (role_list is None or pref > role_list[0]):
                role_list = (pref, name_s, values)

    if mapping is None and role_list is None:
        fail(f"{where}: no sheet with a 'Standard Role' list or 'Position Name' + 'Standard Role' columns was found "
             f"(sheets: {list(sheets)}).")
    if role_list is not None:
        sr = role_list[2]
        if len(sheets) > 1:
            print(f"mapping: standard roles from sheet '{role_list[1]}'")
    else:
        sr = mapping[2]["Standard Role"].dropna().astype(str).str.strip().drop_duplicates()
        sr = sr[sr.ne("")]
        warn(f"{where}: no separate Standard Role list - using the {len(sr)} roles named in the mapping")
    if mapping is None:
        warn(f"{where}: no Position Name -> Standard Role sheet - positions are standardized by the built-in rules only")
    elif len(sheets) > 1:
        print(f"mapping: position names from sheet '{mapping[1]}'")

    roles, redirects = [], {}
    for r in sr:
        m = re.match(r"^(.*?)\s*\(use\s+(.+?)\)\s*$", r, re.I)
        if m:
            redirects[m.group(1).strip().lower()] = m.group(2).strip()
        else:
            roles.append(r)
    if not roles:
        fail(f"{where}: the Standard Role list is empty.")
    tax = Taxonomy(roles, redirects, {})
    if mapping is not None:
        m = mapping[2]
        m = m[m["Standard Role"].notna() & m["Position Name"].notna()]
        unknown = set()
        for name, label in zip(m["Position Name"], m["Standard Role"]):
            role = tax.canonical(label)
            if role:
                tax.examples[norm_key(name)] = role
            else:
                unknown.add(str(label).strip())
        if unknown:
            warn(f"{where}: {len(unknown)} Standard Role labels in the mapping are not in the role list and are ignored "
                 f"(e.g. {sorted(unknown)[:5]})")
    return tax


# =========================================================================== TRANSFORM - standardized positions
def clean_position_text(raw):
    s = "" if pd.isna(raw) else str(raw)
    s = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", s)
    s = re.sub(r"^\s*\d+\s*[.\-)]+\s*", "", s)
    s = re.sub(r"\bHC?R[\s-]?\d+\b|\bHRC\b", " ", s, flags=re.I)
    s = re.sub(r"#\s*\d+", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -–")


def level_text(raw):
    s = "" if pd.isna(raw) else str(raw)
    s = re.sub(r"\bHC?R[\s-]?\d+\b|\bHRC\b|#\s*\d+", " ", s, flags=re.I)
    s = re.sub(r"^\s*\d+\s*[.\-)]+\s*", "", s)
    s = re.sub(r"[\[\](){}]", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -–")


def rule_base(clean, tax):
    for base, rx in ROLE_RULES:
        if base in tax.parse or base in tax.levels or (base, "") in tax.role_of:
            if re.search(rx, clean, re.I):
                return base
    return None


def explicit_level(clean, base, tax):
    """(level, source) read from the name. Only for roles that have levels in the Standard Role sheet."""
    if base not in tax.levels:
        return "", ""
    t = clean.lower()
    m = re.search(r"\bl(?:evel|vl|v)?\s*[-.]?\s*([1-4])(?!\d)", t)
    if m:
        return f"L{m.group(1)}", "name"
    m = re.search(r"(?<![a-z0-9])(iii|ii|i)\s*$", t)
    if m and len(t) > 3:
        return {"i": "L1", "ii": "L2", "iii": "L3"}[m.group(1)], "name"
    if re.search(r"\bsenior\b|\bsr\.?\b|\bsnr\b", t):
        return "L3", "name_title"
    if re.search(r"\bjunior\b|\bjr\b|\bintern\b|\btrainee\b", t):
        return "L1", "name_title"
    if base in LEVEL_DIGIT_BASES:
        m = re.search(r"(?:engineer|cloud)\s*-?\s*([1-3])\s*$", t)
        if m:
            return f"L{m.group(1)}", "name_digit"
    return "", ""


def skill_family(skill):
    if pd.isna(skill):
        return ""
    toks = [t.strip().lower() for t in str(skill).split(";") if t.strip()]
    scores = {f: sum(any(k in t for k in kws) for t in toks) for f, kws in SKILL_FAMILIES.items()}
    fam, top = max(scores.items(), key=lambda kv: kv[1])
    return fam if top > 0 else ""


def resolve_level(base, explicit, grade, tax, lmap):
    lv = tax.levels.get(base, [])
    if not lv:
        return "", ""
    if explicit:
        if explicit in lv:
            return explicit, None                       # keep caller's source label
        return explicit, "name_level_not_in_sheet"
    if len(lv) == 1:
        return lv[0], "single_level"
    g = grade if pd.notna(grade) else 0
    cand = lmap[(lmap["base"] == base) & lmap["level"].isin(lv)] if len(lmap) else lmap
    if len(cand) >= 2 and cand["rows"].sum() >= 3:
        return cand.iloc[(cand["median_grade"] - g).abs().argsort().iloc[0]]["level"], "grade_inferred"
    idx = 0 if g <= 5 else (1 if g <= 7 else 2)
    return lv[min(idx, len(lv) - 1)], "grade_inferred"


def standardise(df, tax):
    """Adds the Standardized Position and skill-family columns; every original column is left untouched."""
    out = df.copy()
    n = len(out)
    clean = out["Position Name"].map(clean_position_text)
    skill = out["Skill"] if "Skill" in out else pd.Series([np.nan] * n, index=out.index)
    base, lvl, lsrc, psrc, role = [None] * n, [""] * n, [""] * n, [""] * n, [None] * n

    # pass 1: exact lookup in the mapping sheet, then rules
    for i, (raw, c) in enumerate(zip(out["Position Name"], clean)):
        hit = tax.examples.get(norm_key(raw))
        if hit:
            role[i], psrc[i] = hit, "mapping_sheet"
            base[i], lvl[i] = tax.parse[hit]
            lsrc[i] = "mapping_sheet" if lvl[i] else ""
            continue
        b = rule_base(c, tax)
        if b:
            base[i], psrc[i] = b, "rule"
            lvl[i], lsrc[i] = explicit_level(level_text(raw), b, tax)

    # grade <-> level evidence from explicit levels (data-driven level inference for names without a level)
    ev = pd.DataFrame({"base": base, "level": lvl, "src": lsrc, "grade": out["Grade"].values})
    ev = ev[ev["src"].isin(["name", "name_digit", "mapping_sheet"]) & (ev["level"] != "")]
    lmap = (ev.groupby(["base", "level"])["grade"].agg(median_grade="median", min_grade="min", max_grade="max", rows="count")
            .reset_index()) if len(ev) else pd.DataFrame(columns=["base", "level", "median_grade", "min_grade", "max_grade", "rows"])

    # pass 2: level resolution for rule hits; unmapped names stay unmapped
    fam = skill.map(skill_family)
    for i in range(n):
        if base[i] is None:
            psrc[i] = "unmapped"
            continue
        if role[i] is None:
            l, s_ = resolve_level(base[i], lvl[i], out["Grade"].iloc[i], tax, lmap)
            lsrc[i] = lsrc[i] if (s_ is None or lsrc[i] == "name_title") else s_
            lvl[i] = l
            role[i] = tax.role_name(base[i], l)
            if role[i] is None:
                psrc[i] = "unmapped"

    out["Standardized Position"] = [r if r and ps != "unmapped" else UNMAPPED for r, ps in zip(role, psrc)]
    out["Std Role Family"] = [b if b and ps != "unmapped" else "" for b, ps in zip(base, psrc)]
    out["Std Level"] = lvl
    out["Grade Band"] = out["Grade"].map(GRADE_BAND).fillna("B?")
    default = out["Std Role Family"].map(BASE_TO_FAMILY).fillna("")
    out["Std Skill Family"] = np.where(fam != "", fam, default)
    return out


# =========================================================================== TRANSFORM - dates / availability
LINE_KEY = ["Project ID", "Quote Line Number", "Position Name", "Grade", "Resource"]


def fy_label(ts):
    fy = ts.year if ts.month >= 4 else ts.year - 1
    return f"FY{fy % 100:02d}"


def add_line_id(df):
    d = df.copy()
    for c in LINE_KEY:
        if c not in d.columns:
            d[c] = np.nan
    d["Slot"] = d.groupby(LINE_KEY + ["Fiscal Year"], dropna=False).cumcount()
    d["Line ID"] = d[LINE_KEY].astype("string").fillna("").agg("|".join, axis=1) + "|" + d["Slot"].astype(str)
    return d


def month_columns(df):
    return [m for m in FY_MONTHS if m in df.columns]


def to_long(d):
    """One row per (line, month) with real 1st-of-month dates. FY26: Apr-Dec 2026, Jan-Mar 2027."""
    d = d.copy()
    d["FY Start Year"] = 2000 + d["Fiscal Year"].str.extract(r"(\d+)")[0].astype(int)
    long = d.melt(id_vars=["Line ID", "FY Start Year"], value_vars=month_columns(d), var_name="Mon", value_name="MDs")
    yr = long["FY Start Year"] + np.where(long["Mon"].isin(["Jan", "Feb", "Mar"]), 1, 0)
    long["Month"] = pd.to_datetime(dict(year=yr, month=long["Mon"].map(MONTH_NUM), day=1))
    long["MDs"] = long["MDs"].fillna(0.0)
    return long.groupby(["Line ID", "Month"], as_index=False)["MDs"].sum()


def line_table(df):
    """One row per line: Start (first month with MDs), Last Active Month, Available From (month after the last MDs)."""
    d = add_line_id(df)
    long = to_long(d)
    act = long[long["MDs"] > 0].sort_values(["Line ID", "Month"])
    g = act.groupby("Line ID")
    s = pd.DataFrame({"Start Date": g["Month"].min(), "Last Active Month": g["Month"].max(),
                      "Avg MDs": g["MDs"].mean(), "Start MDs": g["MDs"].first(), "Last MDs": g["MDs"].last()})
    s["Covered To"] = long.groupby("Line ID")["Month"].max()
    s["Available From"] = s["Last Active Month"] + pd.DateOffset(months=1)
    s["Availability Basis"] = np.where(s["Last Active Month"] < s["Covered To"], "CONFIRMED_END", "FY_END_PLAN")
    s["Start FY"] = s["Start Date"].map(fy_label)
    s["Available From FY"] = s["Available From"].map(fy_label)
    meta = d.drop_duplicates("Line ID").set_index("Line ID")
    return d, meta.join(s, how="inner")


# =========================================================================== MATCH
def month_diff(a, b):
    return (b.year - a.year) * 12 + (b.month - a.month)


def skill_similarity(sup, dem):
    corpus = pd.concat([sup["Skill"].fillna(""), dem["Skill"].fillna("")]).astype(str).str.lower() \
        .str.replace(r"[^a-z0-9+#/ ]", " ", regex=True)
    if corpus.str.strip().eq("").all():
        return np.zeros((len(sup), len(dem))), np.zeros(len(sup), bool), np.zeros(len(dem), bool)
    vec = TfidfVectorizer(ngram_range=(1, 2), stop_words="english").fit(corpus)
    sv, dv = vec.transform(corpus.iloc[:len(sup)]), vec.transform(corpus.iloc[len(sup):])
    return cosine_similarity(sv, dv), sup["Skill"].notna().to_numpy(), dem["Skill"].notna().to_numpy()


def run_match(sup, dem, grade_tol, max_idle, skill_thr):
    sup, dem = sup.reset_index(), dem.reset_index()
    for frame in (sup, dem):
        if "Skill" not in frame.columns:
            frame["Skill"] = np.nan
    ns, nd = len(sup), len(dem)
    sim, s_has, d_has = skill_similarity(sup, dem)
    gap = np.array([[month_diff(a, b) for b in dem["Start Date"]] for a in sup["Available From"]]) if ns and nd else np.zeros((ns, nd), int)
    gd = np.abs(sup["Grade"].to_numpy(float)[:, None] - dem["Grade"].to_numpy(float)[None, :])
    same = sup["Line ID"].to_numpy()[:, None] == dem["Line ID"].to_numpy()[None, :]
    sp, dp = sup["Standardized Position"].to_numpy(), dem["Standardized Position"].to_numpy()
    s_pos_ok = (sup["Standardized Position"] != UNMAPPED).to_numpy()
    d_pos_ok = (dem["Standardized Position"] != UNMAPPED).to_numpy()
    pos_eq = (sp[:, None] == dp[None, :]) & s_pos_ok[:, None] & d_pos_ok[None, :]
    sf, df_ = sup["Std Skill Family"].to_numpy(), dem["Std Skill Family"].to_numpy()
    fam_eq = (sf[:, None] == df_[None, :]) & (sf[:, None] != "")
    both = s_has[:, None] & d_has[None, :]
    skill_score = np.where(both, sim, np.where(fam_eq, 1.0, 0.0))
    skill_ok = fam_eq & (~both | (sim >= skill_thr))
    fy_pen = (sup["Availability Basis"].to_numpy() == "FY_END_PLAN").astype(float)[:, None] * 0.5
    free_ok = (~same) & (gap >= 0) & (gap <= max_idle)              # availability gate: free on/before start

    tiers = [("T1", pos_eq, 0), ("T2", skill_ok, 0), ("T3", pos_eq, grade_tol), ("T4", skill_ok, grade_tol)]
    sfree, dfree, pairs = np.ones(ns, bool), np.ones(nd, bool), []
    for tid, rule, tol in tiers:
        feas = free_ok & rule & (gd <= tol) & sfree[:, None] & dfree[None, :]
        if not feas.any():
            continue
        cost = np.where(feas, gap + 2.0 * gd + (1.0 - skill_score) + fy_pen, BIG)
        r, c = linear_sum_assignment(cost)
        for i, j in zip(r, c):
            if cost[i, j] < BIG:
                pairs.append((i, j, tid, int(gap[i, j])))
                sfree[i], dfree[j] = False, False
    return sup, dem, pairs, sfree, dfree


# =========================================================================== LOAD - the lean output tables
S_COLS = ["Project ID", "Quote Line Number", "Customer", "Project Name", "Resource", "Position Name", "Standardized Position",
          "Grade", "Std Skill Family", "Skill", "Last Active Month", "Available From", "Available From FY",
          "Availability Basis", "Availability Status", "Last MDs"]
D_COLS = ["Project ID", "Quote Line Number", "Customer", "Project Name", "Position Name", "Standardized Position",
          "Grade", "Std Skill Family", "Skill", "Start Date", "Start FY", "Start MDs"]

# Columns of the clean files: the original input columns, then these
CLEAN_EXTRA_COLS = ["Standardized Position", "Std Role Family", "Std Level", "Grade Band", "Std Skill Family",
                    "Start Date", "Last Active Month", "Available From", "Available From FY", "Availability Basis"]
# Input columns that hold dates, shown as real dates in the clean files ("Created On" is an Excel day number in the CSVs)
INPUT_DATE_COLS = ["Created On", "Modified Date"]


def assemble(sup, dem, pairs, sfree, dfree):
    """One row per match, then the unmatched resources, then the unmatched quote lines."""
    sp = sup.reindex(columns=S_COLS).add_prefix("PRP | ")
    dp = dem.reindex(columns=D_COLS).add_prefix("QL | ")
    for p in (sp, dp):
        for c in p.columns:
            if c.endswith("| Skill"):
                p[c] = p[c].astype("string").str.slice(0, 160)
    rows = []
    # matches in tier order (best evidence first), then by start date and waiting time - the order the old output had
    for i, j, tid, g in sorted(pairs, key=lambda p: (p[2], dem.at[p[1], "Start Date"], p[3])):
        r = {"Match Status": "MATCHED"}
        r.update(sp.iloc[i].to_dict())
        r.update(dp.iloc[j].to_dict())
        rows.append(r)
    for i in np.where(sfree)[0]:
        r = {"Match Status": "UNMATCHED RESOURCE"}
        r.update(sp.iloc[i].to_dict())
        rows.append(r)
    for j in np.where(dfree)[0]:
        r = {"Match Status": "UNMATCHED PROJECT"}
        r.update(dp.iloc[j].to_dict())
        rows.append(r)
    return pd.DataFrame(rows, columns=["Match Status"] + list(sp.columns) + list(dp.columns))


def clean_output(df, lines):
    """The input rows with their standardized position and dates; internal helper columns left out."""
    joined = df.join(lines[["Start Date", "Last Active Month", "Available From", "Available From FY",
                            "Availability Basis"]], on="Line ID")
    original = [c for c in df.columns if c not in CLEAN_EXTRA_COLS and c not in ("Line ID", "Slot")]
    out = joined[original + CLEAN_EXTRA_COLS].copy()
    for c in INPUT_DATE_COLS:
        if c in out.columns:
            out[c] = to_date(out[c])
    return out


def to_date(s: pd.Series) -> pd.Series:
    """Excel day numbers (45711) or date text ('09/15/2026 17:05:08') -> dates; anything else is left as it was."""
    num = pd.to_numeric(s, errors="coerce")
    from_serial = pd.to_datetime(num, unit="D", origin="1899-12-30", errors="coerce")
    from_text = pd.to_datetime(s.where(num.isna()), errors="coerce", format="mixed")
    parsed = from_serial.fillna(from_text).dt.normalize()
    return parsed if parsed.notna().sum() >= s.notna().sum() * 0.9 else s


def write_xlsx(path: Path, sheets: dict[str, pd.DataFrame]):
    """Writes the sheets with a coloured header, filters, frozen top row, fitted widths and DD-MM-YY dates."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    status_fill = {"MATCHED": "E2F0D9", "UNMATCHED RESOURCE": "FFF2CC", "UNMATCHED PROJECT": "FCE4D6"}
    try:
        writer = pd.ExcelWriter(path, engine="openpyxl", datetime_format=DATE_FORMAT, date_format=DATE_FORMAT)
    except PermissionError:
        sys.exit(f"Cannot write {path.name} - close it in Excel and run again.")
    try:
        with writer:
            for name, df in sheets.items():
                df = finite(df)
                df.to_excel(writer, sheet_name=name, index=False)
                ws = writer.sheets[name]
                ws.freeze_panes = "A2"
                if len(df) and name != "Summary":
                    ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                for i, col in enumerate(df.columns, start=1):
                    sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200) if pd.notna(v)]
                    ws.column_dimensions[get_column_letter(i)].width = min(max(10, max(sample) + 2), 48)
                    if pd.api.types.is_datetime64_any_dtype(df[col]):
                        ws.column_dimensions[get_column_letter(i)].width = 11
                        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                            cell.number_format = DATE_FORMAT
                if "Match Status" in df.columns and df.columns[0] == "Match Status":
                    for (cell,) in ws.iter_rows(min_row=2, min_col=1, max_col=1):
                        if cell.value in status_fill:
                            cell.fill = PatternFill("solid", fgColor=status_fill[cell.value])
    except PermissionError:
        sys.exit(f"Cannot write {path.name} - close it in Excel and run again.")


# =========================================================================== main
def planning_month(args, prp: pd.DataFrame) -> pd.Timestamp:
    """--as-of, else the month of the latest 'Modified Date' in the resource plan, else the current month."""
    if args.as_of:
        ts = pd.Timestamp(args.as_of)
    elif "Modified Date" in prp.columns and to_date(prp["Modified Date"]).dropna().size:
        ts = pd.to_datetime(to_date(prp["Modified Date"]), errors="coerce").max()
    else:
        ts = pd.Timestamp(date.today())
    return pd.Timestamp(year=ts.year, month=ts.month, day=1)


def main():
    p = argparse.ArgumentParser(description="Match resources rolling off the resource plan to starting quote lines.")
    p.add_argument("--resource-plan", default=None, help="default: a file named like *resource*/*rp*/*plan* in this folder")
    p.add_argument("--quoteline", default=None, help="default: a file named like *quote*/*ql* in this folder")
    p.add_argument("--mapping", default=None, help="Excel/CSV file with a Standard Role list and/or Position Name -> "
                                                   "Standard Role columns (default: a file named like *mapping* in this folder)")
    p.add_argument("--out-dir", default=None, help="write the four output files straight into this folder (default: "
                                                   f"'{OUTPUT_ROOT.name}/{LATEST_DIR}' + '{OUTPUT_ROOT.name}/{HISTORY_DIR}/<run time>')")
    p.add_argument("--as-of", default=None, help="planning month, e.g. 2026-09-01 (default: month of the latest "
                                                 "'Modified Date' in the resource plan, else the current month)")
    p.add_argument("--horizon", type=int, default=12, help="rolling window in months")
    p.add_argument("--grade-tol", type=int, default=1, help="grade +/- allowed in tiers T3/T4 (0 = off)")
    p.add_argument("--max-idle", type=int, default=6, help="max months a free resource may wait for the line start")
    p.add_argument("--skill-thr", type=float, default=0.15, help="min skill similarity when both sides list skills")
    p.add_argument("--confirmed-only", action="store_true", help="ignore lines that simply run to fiscal-year end")
    a = p.parse_args()
    a.mapping = str(locate_input(a.mapping, DEFAULT_MAPPING, "mapping"))
    a.resource_plan = str(locate_input(a.resource_plan, DEFAULT_RESOURCE_PLAN, "resource plan"))
    a.quoteline = str(locate_input(a.quoteline, DEFAULT_QUOTELINE, "quote line"))

    # EXTRACT
    tax = load_taxonomy(a.mapping)
    prp_raw = read_table(Path(a.resource_plan), "resource plan")
    ql_raw = read_table(Path(a.quoteline), "quote line")
    as_of = planning_month(a, prp_raw)

    # TRANSFORM
    prp, ql = standardise(prp_raw, tax), standardise(ql_raw, tax)
    prp_d, prp_lines = line_table(prp)
    ql_d, ql_lines = line_table(ql)
    prp["Line ID"], ql["Line ID"] = prp_d["Line ID"].values, ql_d["Line ID"].values

    # only resources free inside the window, and only lines starting inside it
    end_of_window = as_of + pd.DateOffset(months=a.horizon)
    s = prp_lines[prp_lines["Available From"] < end_of_window].copy()
    if a.confirmed_only:
        s = s[s["Availability Basis"] == "CONFIRMED_END"]
    s["Availability Status"] = np.select(
        [s["Available From"] <= as_of, s["Available From"] < as_of + pd.DateOffset(months=6)],
        ["Available now", "Rolling off 0-6m"], "Rolling off 7-12m")
    d = ql_lines[(ql_lines["Start Date"] >= as_of) & (ql_lines["Start Date"] < end_of_window)].copy()
    if s.empty:
        warn(f"no resource becomes free between {as_of:%d-%m-%y} and {end_of_window:%d-%m-%y} - check --as-of")
    if d.empty:
        warn(f"no quote line starts between {as_of:%d-%m-%y} and {end_of_window:%d-%m-%y} - check --as-of")

    # MATCH
    sup, dem, pairs, sfree, dfree = run_match(s, d, a.grade_tol, a.max_idle, a.skill_thr)
    res = assemble(sup, dem, pairs, sfree, dfree)

    # LOAD
    status = res["Match Status"]
    lean = lambda frame: frame.dropna(axis=1, how="all")          # e.g. no QL columns on unmatched-resource rows
    unmatched_resources = lean(res[status == "UNMATCHED RESOURCE"])
    summary = pd.DataFrame({"Item": [
        "Planning date (as-of)", "Rolling window (months)", "Resources available in window", "Quote lines starting in window",
        "MATCHED", "UNMATCHED RESOURCE", "UNMATCHED PROJECT", "Grade tolerance (+/-)", "Max idle months",
        "Position names not standardized (resource plan / quote line rows)", "Resource plan file", "Quote line file",
        "Mapping file", "Run on"],
        "Value": [as_of.strftime("%d-%m-%y"), a.horizon, len(sup), len(dem),
                  int((status == "MATCHED").sum()), int((status == "UNMATCHED RESOURCE").sum()),
                  int((status == "UNMATCHED PROJECT").sum()), a.grade_tol, a.max_idle,
                  f"{int((prp['Standardized Position'] == UNMAPPED).sum())} / {int((ql['Standardized Position'] == UNMAPPED).sum())}",
                  Path(a.resource_plan).name, Path(a.quoteline).name, Path(a.mapping).name,
                  date.today().strftime("%d-%m-%y")]})

    # without --out-dir each run gets its own Historical Matches folder; Latest is refreshed from it once all four exist
    run_id = datetime.now().strftime(RUN_ID_FORMAT)
    out = Path(a.out_dir) if a.out_dir else OUTPUT_ROOT / HISTORY_DIR / run_id
    out.mkdir(parents=True, exist_ok=True)
    write_xlsx(out / OUT_MATCHING, {
        "Summary": summary,
        "Matching": res,
        "Matched": lean(res[status == "MATCHED"]),
        "Unmatched Resources": unmatched_resources,
        "Unmatched Projects": lean(res[status == "UNMATCHED PROJECT"]),
    })
    write_xlsx(out / OUT_CLEAN_RP, {"Clean Resource Plan": clean_output(prp, prp_lines)})
    write_xlsx(out / OUT_CLEAN_QL, {"Clean Quoteline": clean_output(ql, ql_lines)})
    write_xlsx(out / OUT_UNMATCHED, {"Unmatched Resources": unmatched_resources})

    outputs = (OUT_MATCHING, OUT_CLEAN_RP, OUT_CLEAN_QL, OUT_UNMATCHED)
    print(summary.to_string(index=False))
    print("\nSaved to", out.resolve())
    for name in outputs:
        print("  ", name)
    if not a.out_dir:
        publish_latest(out, outputs)


def publish_latest(run_dir: Path, names):
    """Copy this run's files into Output Files/Latest, replacing the previous run's. A file open in Excel is skipped
    with a warning - the run is still kept in Historical Matches."""
    latest = OUTPUT_ROOT / LATEST_DIR
    latest.mkdir(parents=True, exist_ok=True)
    for name in names:
        try:
            shutil.copy2(run_dir / name, latest / name)
        except PermissionError:
            warn(f"{LATEST_DIR}/{name} is open in Excel - not replaced (this run is in {HISTORY_DIR}/{run_dir.name})")
    print("Latest results in", latest.resolve())


if __name__ == "__main__":
    main()
