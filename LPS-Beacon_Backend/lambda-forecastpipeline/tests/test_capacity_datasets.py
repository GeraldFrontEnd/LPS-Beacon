"""Tests for the capacity-plan inputs the pipeline cleans: the BU UR report and the MSPO staff list
(pipeline/clean.py: clean_bu_ur, clean_staff_list). Both mimic the real exports' layout."""
from __future__ import annotations

import datetime as dt
import io
import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from test_pipeline import Base, s3_event  # noqa: E402

import pandas as pd  # noqa: E402

from pipeline import capacity, clean, insights, plan_settings  # noqa: E402
from test_review import sfdc_xlsx  # noqa: E402

UR_HEADER = ["BU UR Tgt", "Tier & Grade", "Grade Group", "Grade", "Report Month", "Staff Name", "IT Code", "Tier",
             "Billable", "Internal", "Presales", "Admin", "Bench", "Training", "Missing Manday", "Leave",
             "Billable Adj", "Productive Adj", "Submitted Manday", "Standard Working days (Current month)",
             "Billable UR"]


def ur_xlsx() -> bytes:
    """Like the real report: subtotals in row 1 (numbers in I, T and U), the header in row 2, #REF! lookups."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append([None] * 8 + [40.0] + [None] * 10 + [62, 0.645])
    ws.append(UR_HEADER)
    rows = [  # grade, month, name, IT code, billable, working days
        (3, dt.datetime(2026, 4, 15), "Person A", "a1", 18, 21),
        (3, dt.datetime(2026, 4, 15), "Person B", "b1", 21, 21),
        (7, dt.datetime(2026, 4, 15), "Person C", "c1", 10, 21),
        (None, dt.datetime(2026, 4, 15), "Person D", "d1", 0, 21),   # grade lookup broken
        (3, dt.datetime(2026, 5, 15), "Person A", "a1", 19, 19),
    ]
    for grade, month, name, code, billable, days in rows:
        ws.append(["#REF!", f"Tier B - {grade}", "#REF!", grade, month, name, code, "Tier B", billable,
                   0, 0, 0, 0, 0, 0, 0, 0, 0, billable, days, billable / days])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def staff_xlsx() -> bytes:
    """Like the Dynamics export: an empty first row, the header on row 2, and a hidden metadata sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "MSPO SG Staff List"
    ws.append([])
    ws.append(["Name", "IT Code", "Grade", "Join Date", "Last Employment Date", "Cost Center Code",
               "Working City", "Employee Type", "Billable Status", "CCC Billable Status​"])
    people = [("Employee-A", "a1", "3", "Regular Employee"), ("Employee-B", "b1", "3", "Secondment"),
              ("Employee-C", "c1", "7", "Supplemental-Other"), ("Employee-E", "e1", "", "Regular Employee")]
    for name, code, grade, kind in people:
        ws.append([name, code, grade, dt.datetime(2024, 1, 1), None, "SG14000012", "SG - Singapore", kind,
                   "Tier B", "Billable"])
    hidden = wb.create_sheet("hiddenSheet")
    hidden.append(["bookableresource:abc==:name=Name&lvo_itcode=IT%20Code", None])
    hidden.append(["Regular Employee", "Secondment"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class UrReportTests(Base):
    def test_detected_below_the_subtotal_row(self):
        dataset, sheets = clean.detect(clean.read_any(ur_xlsx(), ".xlsx"), "BU_UR.xlsx")
        self.assertEqual(dataset, "bu_ur")
        self.assertIn("Report Month", sheets["main"].columns)

    def test_billable_ur_per_month_and_band(self):
        _, sheets = clean.detect(clean.read_any(ur_xlsx(), ".xlsx"))
        tables, warnings = clean.clean_bu_ur(sheets["main"])
        self.assertNotIn("Staff Name", tables["latest"].columns, "names are dropped")
        m = tables["monthly"]
        overall = m[m["Band"].isna()].set_index("Month")["billableUr"]
        self.assertAlmostEqual(overall["2026-04"], 49 / 84, "all staff, including the row with no grade")
        self.assertAlmostEqual(overall["2026-05"], 1.0)
        band = m.dropna(subset=["Band"]).set_index(["Month", "Band"])["billableUr"]
        self.assertAlmostEqual(band[("2026-04", 3)], 39 / 42)
        self.assertAlmostEqual(band[("2026-04", 7)], 10 / 21)
        self.assertEqual(len(band), 3, "the row without a grade is in no band")
        self.assertEqual(tables["working_days"].set_index("Month")["standardWorkingDays"].to_dict(),
                         {"2026-04": 21, "2026-05": 19})
        self.assertTrue(any("1 rows have no grade" in w for w in warnings))

    def test_upload_stores_it(self):
        import lambda_function
        key = self.put("BU_UR.xlsx", ur_xlsx())
        result = lambda_function.lambda_handler(s3_event(key))
        self.assertEqual(result["files"][0]["dataset"], "bu_ur")
        self.assertTrue(self.store.exists("clean/bu_ur/monthly.json.gz"))
        self.assertTrue(self.store.exists("clean/bu_ur/latest.quality.json"))


class StaffListTests(Base):
    def test_headcount_by_band_and_type(self):
        dataset, sheets = clean.detect(clean.read_any(staff_xlsx(), ".xlsx"), "MSPO SG Staff List.xlsx")
        self.assertEqual(dataset, "staff_list")
        tables, warnings = clean.clean_staff_list(sheets["main"])
        self.assertNotIn("Name", tables["latest"].columns, "names are dropped")
        self.assertIn("CCC Billable Status", tables["latest"].columns, "zero-width space removed from the heading")
        self.assertEqual(tables["headcount_by_band"].set_index("Band")["headcount"].to_dict(), {3: 2, 7: 1})
        self.assertEqual(tables["headcount_by_type"].set_index("Employee Type")["headcount"].to_dict(),
                         {"Regular Employee": 2, "Secondment": 1, "Supplemental-Other": 1})
        self.assertTrue(any("1 staff have no grade" in w for w in warnings))


class CurrencyTests(Base):
    def test_band_rates_in_usd_at_the_finance_rate(self):
        fy26 = plan_settings.resolve("FY26")
        self.assertEqual(fy26["sgdPerUsd"], 1.28)
        self.assertEqual(fy26["bandDayRatesUsd"]["b9"], round(1145 / 1.28, 2))
        self.assertEqual(fy26["bandDayRatesGiven"]["currency"], "SGD")
        self.assertEqual(fy26["revenuePerMonthUsd"], 8_000_000)
        self.assertNotIn("basedOn", fy26)
        self.assertEqual(plan_settings.resolve("FY27")["basedOn"], "FY26", "no FY27 settings yet")
        with self.assertRaises(KeyError):
            plan_settings.resolve("FY25")

    def test_rates_file_is_stored_in_usd(self):
        df = pd.DataFrame({"Grade": ["3", "7"], "Daily Rate": ["256", "500"], "Currency": ["SGD", "usd"]})
        rates, warnings = clean.clean_rates(df, dt.date(2026, 10, 1))
        self.assertEqual(rates["Daily Rate"].tolist(), [200.0, 500.0])
        self.assertEqual(set(rates["Currency"]), {"USD"})
        self.assertTrue(any("1 SGD rates converted" in w for w in warnings))

    def test_rates_file_without_a_known_currency_is_rejected(self):
        with self.assertRaises(clean.CleanError):
            clean.clean_rates(pd.DataFrame({"Grade": ["3"], "Daily Rate": ["200"]}))
        with self.assertRaises(clean.CleanError):
            clean.clean_rates(pd.DataFrame({"Grade": ["3"], "Daily Rate": ["200"], "Currency": ["EUR"]}))


class CapacityInputsTests(Base):
    def test_inputs_after_uploading_the_ur_report_and_staff_list(self):
        import lambda_function
        for name, data in (("BU_UR.xlsx", ur_xlsx()), ("MSPO SG Staff List.xlsx", staff_xlsx())):
            lambda_function.lambda_handler(s3_event(self.put(name, data)))
        inputs = self.store.read_json(capacity.KEY)
        self.assertEqual(inputs["currency"], "USD")
        months = inputs["months"]
        self.assertEqual((len(months), months[0]["month"], months[-1]["month"]), (24, "2026-04", "2028-03"))
        self.assertEqual([m["month"] for m in months if m["actual"]], ["2026-04", "2026-05"])
        self.assertAlmostEqual(inputs["billableUr"]["2026-04"]["b3"], round(39 / 42, 6))
        self.assertAlmostEqual(inputs["billableUr"]["2026-04"]["all"], round(49 / 84, 6))
        self.assertEqual(inputs["workingDays"], {"2026-04": 21.0, "2026-05": 19.0})
        self.assertEqual(inputs["startingHeadcount"], {"month": "2026-04", "bands": {"b3": 2, "b7": 1}})
        hb = inputs["headcountBreakdown"]
        self.assertEqual((hb["permanent"], hb["secondment"]),
                         ({"regular": 2, "supplementalOther": 1, "supplementalRetiree": 0}, 1))
        self.assertEqual(inputs["settings"]["FY27"]["basedOn"], "FY26")
        self.assertTrue(any("FY27 has no settings" in n for n in inputs["notes"]))
        self.assertEqual(inputs["sources"]["bu_ur"]["file"], "abc-BU_UR.xlsx")

    def test_on_demand_with_nothing_uploaded(self):
        import lambda_function
        result = lambda_function.lambda_handler({"mode": "capacity_inputs"})
        self.assertTrue(any("no UR report" in n for n in result["notes"]))
        inputs = self.store.read_json(capacity.KEY)
        self.assertIsNone(inputs["startingHeadcount"])
        self.assertFalse(any(m["actual"] for m in inputs["months"]))


class InsightsTests(Base):
    def test_dashboard_insights_from_the_uploaded_files(self):
        import lambda_function
        for name, data in (("BU_UR.xlsx", ur_xlsx()), ("MSPO SG Staff List.xlsx", staff_xlsx()),
                           ("YTD_Pipeline.xlsx", sfdc_xlsx())):
            lambda_function.lambda_handler(s3_event(self.put(name, data)))
        monthly = self.store.read_table("clean/sfdc_pipeline/monthly_revenue.json.gz")
        self.assertEqual(sorted(monthly[monthly["CRM ID"] == "CRM-1"]["Month"]),
                         ["2026-10", "2026-11", "2026-12", "2027-01", "2027-02", "2027-03"])

        data = insights.build(self.store, dt.date(2026, 10, 15))
        pipe = data["pipeline"]
        self.assertEqual(pipe["kpis"]["openCount"], 3, "O4, O3 and O2; Closed Lost is not open")
        self.assertEqual([s["stage"][:2] for s in pipe["byStage"]], ["O2", "O3", "O4"])
        o4 = next(s for s in pipe["byStage"] if s["stage"].startswith("O4"))
        self.assertEqual((o4["tcv"], o4["weightedTcv"]), (100000, 60000), "TCV x probability")
        oct_ = pipe["revenueByMonth"][0]
        self.assertEqual((oct_["month"], oct_["open"], oct_["openWeighted"]), ("2026-10", 10, 6.0))
        self.assertEqual(pipe["topOpen"][0]["crmId"], "CRM-2", "largest weighted TCV first")
        self.assertEqual(pipe["kpis"]["winRate"], 0.0, "one closed in the last 12 months, and it was lost")

        util = data["utilisation"]
        self.assertEqual(util["latestMonth"], "2026-05")
        self.assertEqual(util["overallTarget"], 0.854)
        self.assertEqual(util["nonBillableByMonth"][0]["workingDays"], 84)
        self.assertIsNone(data["resources"], "no review completed, so no roll-off figures")
        self.assertEqual(data["workforce"]["upcomingLeaversTotal"], 0)
        self.assertTrue(self.store.exists(insights.KEY), "written when the files were cleaned")


if __name__ == "__main__":
    unittest.main()
